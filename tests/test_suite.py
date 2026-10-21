"""
Test suite for Catalog Manager.
Run: pytest tests/ -v --tb=short
"""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent.parent))

import pytest
import tempfile
import shutil
from datetime import datetime

from app.models.product import Product, MarginConfig, ProcessingSession
from app.parsers.base_parser import BaseParser
from app.parsers.parser_factory import ParserFactory
from app.database.db_manager import DatabaseManager
from app.services.excel_service import ExcelService


# ─── Fixtures ────────────────────────────────────────────────────────────────

@pytest.fixture
def tmp_dir():
    d = Path(tempfile.mkdtemp())
    yield d
    shutil.rmtree(d, ignore_errors=True)


@pytest.fixture
def db(tmp_dir):
    return DatabaseManager(tmp_dir / "test.db")


@pytest.fixture
def excel_svc(db, tmp_dir):
    return ExcelService(db, tmp_dir / "excels")


@pytest.fixture
def sample_product():
    return Product(
        sku="ACC01",
        brand="ACCESSORIES",
        description="Silicone hair styling bands",
        price_original=1834.29,
        category="ACCESSORIES",
    )


@pytest.fixture
def sample_product_list():
    return [
        Product("SKU001", "BRAND A", "Product 1", 1000.0, "CAT A"),
        Product("SKU002", "BRAND B", "Product 2", 2500.0, "CAT A"),
        Product("SKU003", "BRAND C", "Product 3", 500.0,  "CAT B"),
    ]


# ─── Model tests ──────────────────────────────────────────────────────────────

class TestProduct:
    def test_creation(self, sample_product):
        p = sample_product
        assert p.sku == "ACC01"
        assert p.price_original == 1834.29
        assert p.active is True
        assert p.first_seen is not None

    def test_apply_margin_simple(self, sample_product):
        p = sample_product
        final = p.apply_margin(30.0)
        assert final == pytest.approx(1834.29 * 1.30, rel=1e-3)
        assert p.price_final == final
        assert p.margin_pct == 30.0

    def test_apply_margin_with_iva(self, sample_product):
        p = sample_product
        final = p.apply_margin(30.0, iva_pct=21.0)
        expected = 1834.29 * 1.30 * 1.21
        assert final == pytest.approx(expected, rel=1e-3)

    def test_apply_margin_round_to(self, sample_product):
        p = sample_product
        p.apply_margin(30.0, round_to=100)
        assert p.price_final % 100 == 0

    def test_apply_margin_invalid_suffix_ignored(self, sample_product):
        """C2: invalid suffix must not corrupt the price."""
        p = sample_product
        final = p.apply_margin(30.0, suffix="abc")   # invalid — must be ignored
        assert p.price_suffix == ""
        assert isinstance(final, float)
        assert final == pytest.approx(1834.29 * 1.30, rel=1e-3)

    def test_apply_margin_valid_suffix(self, sample_product):
        """C2: valid suffix (.99) is applied after rounding."""
        p = sample_product
        p.apply_margin(30.0, round_to=100, suffix=".99")
        assert p.price_suffix == ".99"
        # price_final should end in .99
        frac = round(p.price_final % 1, 2)
        assert frac == pytest.approx(0.99, abs=0.01)

    def test_apply_margin_comma_suffix(self, sample_product):
        """Comma-formatted suffix (Argentine style) must also be accepted."""
        p = sample_product
        p.apply_margin(30.0, round_to=100, suffix=",99")
        assert p.price_suffix == ",99"
        frac = round(p.price_final % 1, 2)
        assert frac == pytest.approx(0.99, abs=0.01)


class TestBaseParserPriceParser:
    def test_argentine_format(self):
        assert BaseParser._parse_price("$ 1.484,90") == pytest.approx(1484.90)
        assert BaseParser._parse_price("$ 131.928,30") == pytest.approx(131928.30)
        assert BaseParser._parse_price("$ 1.362,62") == pytest.approx(1362.62)

    def test_simple_format(self):
        assert BaseParser._parse_price("1500.00") == pytest.approx(1500.0)
        assert BaseParser._parse_price("500") == pytest.approx(500.0)

    def test_invalid(self):
        assert BaseParser._parse_price("no price") == 0.0
        assert BaseParser._parse_price("") == 0.0

    def test_normalize_sku(self):
        assert BaseParser._normalize_sku("  acc01  ") == "ACC01"
        assert BaseParser._normalize_sku("acce03") == "ACCE03"


class TestMarginConfig:
    def test_defaults(self):
        mc = MarginConfig()
        assert mc.margin_pct == 30.0
        assert mc.iva_pct == 0.0
        assert mc.round_to_nearest == 0

    def test_custom(self):
        mc = MarginConfig(margin_pct=45.0, iva_pct=21.0, round_to_nearest=100)
        assert mc.margin_pct == 45.0


# ─── Database tests ───────────────────────────────────────────────────────────

class TestDatabaseManager:
    def test_upsert_and_get(self, db, sample_product):
        data = {
            "sku": sample_product.sku,
            "brand": sample_product.brand,
            "description": sample_product.description,
            "category": sample_product.category,
            "price_original": sample_product.price_original,
            "price_final": sample_product.price_final,
            "margin_pct": 30.0,
            "iva_pct": 0.0,
            "active": True,
            "discontinued": False,
            "notes": "",
            "image_path": None,
            "image_confidence": 0.0,
            "first_seen": datetime.now(),
            "last_seen": datetime.now(),
            "processing_id": "TEST001",
            "page_num": 0,
        }
        db.upsert_product(data)
        result = db.get_product("ACC01")
        assert result is not None
        assert result["sku"] == "ACC01"
        assert result["price_original"] == pytest.approx(1834.29)

    def test_active_flag_preserved_on_update(self, db):
        """Critical: once deactivated, a product stays deactivated on re-import."""
        base_data = {
            "sku": "TEST01", "brand": "B", "description": "D",
            "category": "C", "price_original": 100.0, "price_final": 130.0,
            "margin_pct": 30.0, "iva_pct": 0.0, "active": True,
            "discontinued": False, "notes": "", "image_path": None,
            "image_confidence": 0.0, "first_seen": datetime.now(),
            "last_seen": datetime.now(), "processing_id": "T", "page_num": 0,
        }
        db.upsert_product(base_data)
        db.set_product_active("TEST01", False)

        # Re-import same product (simulating new PDF with same SKU)
        reimport_data = {**base_data, "price_original": 110.0}
        db.upsert_product(reimport_data)

        result = db.get_product("TEST01")
        assert result["active"] is False, "Active flag must be preserved after re-import"
        assert result["price_original"] == pytest.approx(110.0), "Price must be updated"

    def test_mark_discontinued(self, db):
        for i in range(3):
            db.upsert_product({
                "sku": f"SKU{i:02d}", "brand": "B", "description": "D",
                "category": "C", "price_original": 100.0, "price_final": 100.0,
                "margin_pct": 0.0, "iva_pct": 0.0, "active": True,
                "discontinued": False, "notes": "", "image_path": None,
                "image_confidence": 0.0, "first_seen": datetime.now(),
                "last_seen": datetime.now(), "processing_id": "T", "page_num": i,
            })
        # Only SKU00 seen in new import
        count = db.mark_discontinued(["SKU00"])
        assert count == 2
        assert db.get_product("SKU01")["discontinued"] is True
        assert db.get_product("SKU00")["discontinued"] is False

    def test_discontinued_flag_preserved_on_upsert(self, db):
        """I5: discontinued flag must NOT be reset when re-importing a known SKU."""
        base = {
            "sku": "DISC01", "brand": "B", "description": "D",
            "category": "C", "price_original": 100.0, "price_final": 130.0,
            "margin_pct": 30.0, "iva_pct": 0.0, "active": True,
            "discontinued": False, "notes": "", "image_path": None,
            "image_confidence": 0.0, "first_seen": datetime.now(),
            "last_seen": datetime.now(), "processing_id": "T", "page_num": 0,
        }
        db.upsert_product(base)
        db.mark_discontinued(["OTHER_SKU"])  # marks DISC01 discontinued
        assert db.get_product("DISC01")["discontinued"] is True

        # Re-import same SKU without discontinued in payload
        reimport = {k: v for k, v in base.items() if k != "discontinued"}
        reimport["price_original"] = 120.0
        db.upsert_product(reimport)

        result = db.get_product("DISC01")
        assert result["discontinued"] is True, "discontinued must survive re-import"
        assert result["price_original"] == pytest.approx(120.0)

    def test_settings(self, db):
        db.set_setting("company_name", "Test Corp")
        assert db.get_setting("company_name") == "Test Corp"
        assert db.get_setting("missing_key", "default") == "default"

        db.set_setting("company_name", "Updated Corp")
        assert db.get_setting("company_name") == "Updated Corp"

    def test_get_all_products(self, db, sample_product_list):
        for p in sample_product_list:
            db.upsert_product({
                "sku": p.sku, "brand": p.brand, "description": p.description,
                "category": p.category, "price_original": p.price_original,
                "price_final": p.price_final, "margin_pct": 0.0, "iva_pct": 0.0,
                "active": True, "discontinued": False, "notes": "",
                "image_path": None, "image_confidence": 0.0,
                "first_seen": datetime.now(), "last_seen": datetime.now(),
                "processing_id": "T", "page_num": 0,
            })
        all_p = db.get_all_products(include_inactive=True)
        assert len(all_p) == 3

        db.set_product_active("SKU002", False)
        active_p = db.get_all_products(include_inactive=False)
        assert len(active_p) == 2


# ─── Parser factory tests ─────────────────────────────────────────────────────

class TestParserFactory:
    def test_list_parsers(self):
        parsers = ParserFactory.list_parsers()
        assert len(parsers) >= 1
        assert any(p["name"] == "Mayoristas I" for p in parsers)

    def test_sample_text_extraction_failure(self, tmp_dir):
        fake_pdf = tmp_dir / "fake.pdf"
        fake_pdf.write_bytes(b"not a pdf")
        result = ParserFactory._extract_sample_text(fake_pdf)
        assert result == ""

    def test_no_match_returns_none(self, tmp_dir):
        """C4: unrecognised PDF must return None, not a blind fallback parser."""
        fake_pdf = tmp_dir / "unknown_supplier.pdf"
        fake_pdf.write_bytes(b"not a pdf")
        parser_cls = ParserFactory.detect(fake_pdf)
        assert parser_cls is None, "ParserFactory must return None for unrecognised PDFs"


# ─── Excel service tests ──────────────────────────────────────────────────────

class TestExcelService:
    def _populate_db(self, db, products):
        for p in products:
            db.upsert_product({
                "sku": p.sku, "brand": p.brand, "description": p.description,
                "category": p.category, "price_original": p.price_original,
                "price_final": p.price_final, "margin_pct": 30.0, "iva_pct": 0.0,
                "active": True, "discontinued": False, "notes": "",
                "image_path": None, "image_confidence": 0.0,
                "first_seen": datetime.now(), "last_seen": datetime.now(),
                "processing_id": "TEST", "page_num": 0,
            })

    def test_generate_excel(self, excel_svc, db, tmp_dir, sample_product_list):
        self._populate_db(db, sample_product_list)
        path = excel_svc.generate_excel()
        assert path.exists()
        assert path.suffix == ".xlsx"
        assert path.stat().st_size > 0
        # Two sheets
        from openpyxl import load_workbook
        wb = load_workbook(str(path), data_only=False)
        assert "Products" in wb.sheetnames, "Products sheet must exist"
        assert "System"   in wb.sheetnames, "System sheet must exist"
        # Final Price is a pre-calculated numeric value (not a formula) — C1 fix
        # Row 1 = staleness banner, Row 2 = headers, Row 3+ = data
        ws = wb["Products"]
        cell_val = ws.cell(3, 9).value
        assert cell_val is not None, "Final Price must not be empty"
        assert isinstance(cell_val, (int, float)), (
            f"Final Price must be a number (not a formula), got: {cell_val!r}"
        )
        assert ws.cell(3, 9).number_format == '"$"#,##0.00'
        # The image column must NOT be on sheet 1 (headers are on row 2)
        headers = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        assert "Image" not in headers, "Image must not be on the Products sheet"

    def test_generate_excel_empty_db_raises(self, excel_svc, db):
        with pytest.raises(ValueError, match="no products"):
            excel_svc.generate_excel()

    def test_reimport_excel(self, excel_svc, db, tmp_dir, sample_product_list):
        self._populate_db(db, sample_product_list)
        # Export
        excel_path = excel_svc.generate_excel()
        # Edit: deactivate SKU001 with openpyxl
        import openpyxl
        wb = openpyxl.load_workbook(str(excel_path))
        ws = wb["Products"]
        # Row 1 = banner, Row 2 = headers, Row 3+ = data
        for row in ws.iter_rows(min_row=3):
            if row[1].value == "SKU001":  # col B = SKU
                row[0].value = "NO"       # col A = Active
        wb.save(str(excel_path))
        # Reimport
        results = excel_svc.reimport_excel(excel_path)
        assert results["processed"] == 3
        assert results["updated"] >= 1
        # SKU001 must now be inactive
        p = db.get_product("SKU001")
        assert p["active"] is False

    def test_reimport_missing_sku_column(self, excel_svc, tmp_dir):
        import pandas as pd
        bad_excel = tmp_dir / "bad.xlsx"
        pd.DataFrame({"Wrong": [1, 2]}).to_excel(str(bad_excel), index=False)
        results = excel_svc.reimport_excel(bad_excel)
        assert len(results["errors"]) > 0


# ─── Processing session tests ─────────────────────────────────────────────────

class TestProcessingSession:
    def test_session_creation(self):
        s = ProcessingSession(pdf_path="test.pdf")
        assert s.session_id  # auto-generated
        assert s.products_found == 0
        assert s.completed_at is None

    def test_session_complete(self):
        s = ProcessingSession()
        s.complete()
        assert s.completed_at is not None
        assert s.duration_seconds >= 0


# ─── Formatting helpers ───────────────────────────────────────────────────────

class TestFormatting:
    def test_format_price_us(self):
        from app.utils.formatting import format_price
        assert format_price(1234.5) == "$1,234.50"
        assert format_price(18.06) == "$18.06"
        assert format_price(1000, decimals=0) == "$1,000"

    def test_parse_number_accepts_us_and_comma_decimal(self):
        from app.utils.formatting import parse_number
        assert parse_number("1,234.56") == pytest.approx(1234.56)
        assert parse_number("$ 18.06") == pytest.approx(18.06)
        assert parse_number("12,5") == pytest.approx(12.5)
        assert parse_number("30") == pytest.approx(30.0)

    def test_parse_number_rejects_garbage(self):
        from app.utils.formatting import parse_number
        with pytest.raises(ValueError):
            parse_number("abc")
        with pytest.raises(ValueError):
            parse_number("")


# ─── Legacy Excel compatibility ───────────────────────────────────────────────

class TestLegacyExcelReimport:
    def test_spanish_headers_still_reimport(self, excel_svc, db, tmp_dir, sample_product_list):
        """Excels exported before the English UI used Spanish headers and SI/NO."""
        TestExcelService()._populate_db(db, sample_product_list)
        import openpyxl
        legacy = tmp_dir / "legacy.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["banner"])
        ws.append(["Activo", "SKU", "Marca", "Descripción", "Categoría",
                   "Precio Original", "Margen %", "IVA %", "Precio Final", "Observaciones"])
        ws.append(["NO", "SKU001", "BRAND A", "Product 1", "CAT A", 1000, 30, 0, 1300, ""])
        ws.append(["SI", "SKU002", "BRAND B", "Product 2", "CAT A", 2500, 30, 0, 3250, ""])
        wb.save(str(legacy))

        results = excel_svc.reimport_excel(legacy)
        assert results["errors"] == []
        assert results["updated"] == 2
        assert db.get_product("SKU001")["active"] is False
        assert db.get_product("SKU002")["active"] is True
