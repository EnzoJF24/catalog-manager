"""
Excel service v3:
  Sheet 1 "Products"  — editable by the user.
                         Final Price is a numeric VALUE (not a formula),
                         computed by apply_margin() and stored in the DB.
                         The user may edit Margin % / Tax % directly; on
                         re-import the price is recalculated with apply_margin().
  Sheet 2 "System"    — internal data: image↔SKU, metadata, configuration.
  Reimport            — recalculates price_final with apply_margin() using the
                         round_to/suffix persisted in the DB (single source of truth).
                         Accepts both English headers and the legacy Spanish ones.
"""
from __future__ import annotations
from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side,
)
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.database.db_manager import DatabaseManager
from app.logging_config.logger import get_logger
from app.config.branding import APP_NAME, APP_VERSION, APP_AUTHOR, EXCEL_SYSTEM_LABEL
from app.config.settings import settings
from app.utils.formatting import EXCEL_PRICE_FORMAT, parse_number

log = get_logger("excel_service")

TRUTHY = {"yes", "y", "true", "1", "active", "x", "si", "sí", "activo"}

# ── Column map for Sheet 1 ────────────────────────────────────────────────────
COL_ACTIVO  = 1
COL_SKU     = 2
COL_MARCA   = 3
COL_DESC    = 4
COL_CAT     = 5
COL_ORIG    = 6
COL_MARGEN  = 7
COL_IVA     = 8
COL_FINAL   = 9   # ← numeric value (pre-calculated), no formula
COL_OBS     = 10

HEADERS_USER = [
    "Active", "SKU", "Brand", "Description", "Category",
    "Cost", "Margin %", "Tax %", "Final Price",
    "Notes",
]

# Re-import accepts the current headers and the legacy Spanish ones, so files
# exported by earlier versions of the app can still be pushed back.
HEADER_ALIASES = {
    "active":      ("Active", "Activo"),
    "sku":         ("SKU",),
    "brand":       ("Brand", "Marca"),
    "description": ("Description", "Descripción"),
    "category":    ("Category", "Categoría"),
    "cost":        ("Cost", "Precio Original"),
    "margin":      ("Margin %", "Margen %"),
    "tax":         ("Tax %", "IVA %"),
    "final":       ("Final Price", "Precio Final"),
    "notes":       ("Notes", "Observaciones"),
}

# ── Palette (mirrors app/ui/theme.py for Excel consistency) ──────────────────
_RED        = "C0392B"
_DARK_BLUE  = "16213E"
_DEEP_BLUE  = "0F1E36"
_LIGHT      = "F5F5F5"
_GREEN_BG   = "E8F8E8"
_RED_BG     = "F8E8E8"
_BLUE_BG    = "EBF5FB"   # Final Price column
_GREY_TAB   = "95A5A6"
_SYS_HDR    = "2C3E50"
_SYS_SEC    = "1A252F"

_BANNER_BG  = "FCF3CF"   # soft yellow — staleness warning banner
_BANNER_TXT = "7D6608"   # dark gold — banner text

_FILL_HEADER   = PatternFill("solid", fgColor=_RED)
_FILL_BANNER   = PatternFill("solid", fgColor=_BANNER_BG)
_FILL_ACTIVE   = PatternFill("solid", fgColor=_GREEN_BG)
_FILL_INACTIVE = PatternFill("solid", fgColor=_RED_BG)
_FILL_ALT      = PatternFill("solid", fgColor=_LIGHT)
_FILL_FINAL    = PatternFill("solid", fgColor=_BLUE_BG)
_FILL_SYS_HDR  = PatternFill("solid", fgColor=_SYS_HDR)
_FILL_SYS_SEC  = PatternFill("solid", fgColor=_SYS_SEC)

_FONT_HEADER = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
_FONT_BANNER = Font(bold=True, name="Calibri", size=11, color=_BANNER_TXT)
_FONT_DATA   = Font(name="Calibri", size=10)
_FONT_BOLD   = Font(bold=True, name="Calibri", size=10)
_FONT_FINAL  = Font(bold=True, name="Calibri", size=10, color="154360")
_FONT_SYS    = Font(name="Consolas", size=9, color="BDC3C7")
_FONT_SYS_H  = Font(bold=True, name="Calibri", size=11, color="FFFFFF")
_FONT_SYS_S  = Font(bold=True, name="Calibri", size=10, color="3498DB")
_FONT_NOTE   = Font(italic=True, name="Calibri", size=9, color="7F8C8D")

_ALIGN_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
_ALIGN_L = Alignment(horizontal="left",   vertical="center", wrap_text=True)
# Data rows have a fixed height: wrapping would cut long descriptions in half
_ALIGN_ROW = Alignment(horizontal="left", vertical="center", wrap_text=False)
_ALIGN_R = Alignment(horizontal="right",  vertical="center")
_BORDER  = Border(
    bottom=Side(style="thin", color="DDDDDD"),
    right=Side(style="thin",  color="DDDDDD"),
)


def _cell(ws, row: int, col: int):
    return ws.cell(row=row, column=col)


def _safe_text(value):
    """Strip control characters openpyxl refuses to write (e.g. NUL from PDF text)."""
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


class ExcelService:

    def __init__(self, db: DatabaseManager, output_dir: Path):
        self.db = db
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    # ── Generation ────────────────────────────────────────────────────────────

    def generate_excel(
        self,
        filename: Optional[str] = None,
        active_only: bool = False,
        include_discontinued: bool = False,
        round_to: int = 0,
        suffix: str = "",
    ) -> Path:
        rows = self.db.get_all_products(include_inactive=True)
        if not rows:
            raise ValueError("There are no products in the database.")

        if active_only:
            rows = [r for r in rows if r.get("active")]
        if not include_discontinued:
            rows = [r for r in rows if not r.get("discontinued")]
        if not rows:
            raise ValueError("No products match the filter criteria.")

        rows.sort(key=lambda r: (r.get("category") or "", r.get("sku") or ""))

        if filename:
            out_path = self.output_dir / filename
        else:
            base     = f"catalog_{datetime.now().strftime('%Y-%m-%d')}"
            out_path = self.output_dir / f"{base}.xlsx"
            # If a file with that date already exists, append _2, _3, … to avoid overwrite
            n = 2
            while out_path.exists():
                out_path = self.output_dir / f"{base}_{n}.xlsx"
                n += 1

        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Products"
        self._write_products_sheet(ws1, rows, self._header_fill())

        ws2 = wb.create_sheet("System")
        ws2.sheet_properties.tabColor = _GREY_TAB
        self._write_system_sheet(ws2, rows, round_to, suffix)

        wb.save(str(out_path))
        log.info(f"Excel generated: {out_path} ({len(rows)} rows)")
        return out_path

    # ── Sheet 1 ───────────────────────────────────────────────────────────────

    def _header_fill(self) -> PatternFill:
        """Header colour follows the catalog's primary colour (Catalog tab settings)."""
        value = (self.db.get_setting("primary_color", "") or "").strip().lstrip("#")
        if len(value) == 6:
            try:
                int(value, 16)
                return PatternFill("solid", fgColor=value.upper())
            except ValueError:
                pass
        return _FILL_HEADER

    def _write_products_sheet(self, ws, rows: list[dict], header_fill: PatternFill = _FILL_HEADER):
        # ── Row 1: staleness banner (merged across all columns) ──────────────
        gen_str = datetime.now().strftime("%m/%d/%Y at %H:%M")
        banner  = _cell(ws, 1, 1)
        banner.value = (
            f"⚠  Catalog generated on {gen_str}. "
            "Changes made in the app after that are NOT reflected here. "
            "For up-to-date data, export a new Excel from the application."
        )
        banner.fill      = _FILL_BANNER
        banner.font      = _FONT_BANNER
        banner.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.merge_cells(
            start_row=1, start_column=1, end_row=1, end_column=COL_OBS
        )
        ws.row_dimensions[1].height = 38

        # ── Row 2: headers ────────────────────────────────────────────────────
        for col, h in enumerate(HEADERS_USER, 1):
            c            = _cell(ws, 2, col)
            c.value      = h
            c.fill       = header_fill
            c.font       = _FONT_HEADER
            c.alignment  = _ALIGN_C
            c.border     = _BORDER

        # Dropdown for "Active" column
        dv = DataValidation(type="list", formula1='"YES,NO"', allow_blank=False)
        dv.error       = "Use YES or NO"
        dv.errorTitle  = "Invalid value"
        dv.prompt      = "YES = active, NO = inactive"
        dv.promptTitle = "Product status"
        ws.add_data_validation(dv)

        # ── Row 3+: data ──────────────────────────────────────────────────────
        for row_i, r in enumerate(rows, start=3):
            is_active = r.get("active", True)

            # A — Active
            c           = _cell(ws, row_i, COL_ACTIVO)
            c.value     = "YES" if is_active else "NO"
            c.fill      = _FILL_ACTIVE if is_active else _FILL_INACTIVE
            c.font      = Font(bold=True, name="Calibri", size=10,
                               color="1A6B3C" if is_active else "8B0000")
            c.alignment = _ALIGN_C
            c.border    = _BORDER
            dv.add(c)

            # B-E — text
            for col, val in [
                (COL_SKU,   r.get("sku", "")),
                (COL_MARCA, r.get("brand") or ""),
                (COL_DESC,  r.get("description") or ""),
                (COL_CAT,   r.get("category") or ""),
            ]:
                c           = _cell(ws, row_i, col)
                c.value     = _safe_text(val)
                c.font      = _FONT_DATA
                c.alignment = _ALIGN_ROW
                c.border    = _BORDER
                if row_i % 2 == 0 and col != COL_SKU:
                    c.fill = _FILL_ALT

            # F — Cost
            c               = _cell(ws, row_i, COL_ORIG)
            c.value         = round(r.get("price_original", 0.0) or 0.0, 2)
            c.font          = _FONT_DATA
            c.alignment     = _ALIGN_R
            c.border        = _BORDER
            c.number_format = EXCEL_PRICE_FORMAT

            # G — Margin %
            c               = _cell(ws, row_i, COL_MARGEN)
            c.value         = round(r.get("margin_pct", 0.0) or 0.0, 2)
            c.font          = _FONT_DATA
            c.alignment     = _ALIGN_R
            c.border        = _BORDER
            c.number_format = '0.00"%"'

            # H — Tax %
            c               = _cell(ws, row_i, COL_IVA)
            c.value         = round(r.get("iva_pct", 0.0) or 0.0, 2)
            c.font          = _FONT_DATA
            c.alignment     = _ALIGN_R
            c.border        = _BORDER
            c.number_format = '0.00"%"'

            # I — Final Price (numeric, pre-calculated — no live formula)
            c               = _cell(ws, row_i, COL_FINAL)
            c.value         = round(r.get("price_final", 0.0) or 0.0, 2)
            c.font          = _FONT_FINAL
            c.fill          = _FILL_FINAL
            c.alignment     = _ALIGN_R
            c.border        = _BORDER
            c.number_format = EXCEL_PRICE_FORMAT

            # J — Notes
            c           = _cell(ws, row_i, COL_OBS)
            c.value     = _safe_text(r.get("notes") or "")
            c.font      = _FONT_DATA
            c.alignment = _ALIGN_ROW
            c.border    = _BORDER

        # Column widths
        for col, w in {
            # Header widths leave room for the auto-filter dropdown arrow
            COL_ACTIVO: 11, COL_SKU: 15, COL_MARCA: 18,
            COL_DESC: 88, COL_CAT: 18, COL_ORIG: 13,
            COL_MARGEN: 14, COL_IVA: 11, COL_FINAL: 17, COL_OBS: 36,
        }.items():
            ws.column_dimensions[get_column_letter(col)].width = w

        # Explanatory note below the data (data starts at row 3 now)
        help_row = len(rows) + 4
        note = _cell(ws, help_row, COL_FINAL - 2)
        note.value     = (
            "ℹ  Final Price is the value calculated by the system. "
            "To recalculate it, use 'Re-apply margin' in Tools."
        )
        note.font      = _FONT_NOTE
        note.alignment = _ALIGN_L
        ws.merge_cells(
            start_row=help_row, start_column=COL_FINAL - 2,
            end_row=help_row,   end_column=COL_OBS
        )

        # Freeze under headers (row 2), filter on header row, set row heights
        ws.freeze_panes    = "A3"
        ws.auto_filter.ref = f"A2:{get_column_letter(COL_OBS)}2"
        ws.row_dimensions[2].height = 22
        for r_i in range(3, len(rows) + 3):
            ws.row_dimensions[r_i].height = 20

    # ── Sheet 2 ───────────────────────────────────────────────────────────────

    def _write_system_sheet(self, ws, rows: list[dict], round_to: int, suffix: str):
        def _section_title(row: int, text: str):
            c           = _cell(ws, row, 1)
            c.value     = f"  {text}"
            c.fill      = _FILL_SYS_SEC
            c.font      = _FONT_SYS_S
            c.alignment = _ALIGN_L
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)

        def _kv(row: int, key: str, val):
            k           = _cell(ws, row, 1)
            k.value     = key
            k.font      = Font(bold=True, name="Calibri", size=10, color="7F8C8D")
            k.alignment = _ALIGN_R
            v           = _cell(ws, row, 2)
            v.value     = str(val)
            v.font      = Font(name="Calibri", size=10, color="FFFFFF")
            v.alignment = _ALIGN_L

        # Global header
        c           = _cell(ws, 1, 1)
        c.value     = EXCEL_SYSTEM_LABEL
        c.fill      = _FILL_SYS_HDR
        c.font      = _FONT_SYS_H
        c.alignment = _ALIGN_L
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
        ws.row_dimensions[1].height = 24

        c2           = _cell(ws, 2, 1)
        c2.value     = (
            f"  Generated by {APP_NAME} v{APP_VERSION}  ·  {APP_AUTHOR}  "
            "— This sheet is for internal system use. Do not edit."
        )
        c2.font      = _FONT_NOTE
        c2.alignment = _ALIGN_L
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=8)

        # Generation metadata
        _section_title(4, "GENERATION METADATA")
        _kv(5, "Generated on:",        datetime.now().strftime("%m/%d/%Y %H:%M:%S"))
        _kv(6, "Total products:",      len(rows))
        sessions = self.db.get_sessions(limit=1)
        if sessions:
            s = sessions[0]
            _kv(7, "Last session ID:",   s.get("session_id", "—"))
            _kv(8, "Supplier:",          s.get("supplier_name", "—"))
            _kv(9, "Parser:",            s.get("parser_name", "—"))

        # Applied config
        _section_title(11, "APPLIED CONFIGURATION")
        sample = rows[0] if rows else {}
        _kv(12, "Margin %:",         f"{sample.get('margin_pct', 0):.2f}")
        _kv(13, "Tax %:",            f"{sample.get('iva_pct', 0):.2f}")
        _kv(14, "Round up to:",      round_to if round_to else "No rounding")
        _kv(15, "Price ending:",     suffix if suffix else "None")
        _kv(16, "Final Price:",      "Pre-calculated by the system (not a formula)")

        # SKU ↔ image table
        _section_title(18, "IMAGE ↔ PRODUCT TABLE")
        img_headers = [
            "SKU", "Image Path", "Image Exists",
            "Brand", "Category", "First Seen",
            "Last Import", "Session ID",
        ]
        for col, h in enumerate(img_headers, 1):
            c           = _cell(ws, 19, col)
            c.value     = h
            c.fill      = _FILL_SYS_HDR
            c.font      = _FONT_SYS_H
            c.alignment = _ALIGN_C
            c.border    = _BORDER

        img_alt = PatternFill("solid", fgColor="1A252F")
        for row_i, r in enumerate(rows, start=20):
            img_path = r.get("image_path") or ""
            if not img_path:
                img_ok = "— no image"
            elif settings.resolve_path(img_path).exists():
                img_ok = "✓"
            else:
                img_ok = "⚠ file not found"

            vals = [
                r.get("sku", ""),
                img_path,
                img_ok,
                r.get("brand", ""),
                r.get("category", ""),
                str(r.get("first_seen", ""))[:16],
                str(r.get("last_seen",  ""))[:16],
                r.get("processing_id", ""),
            ]
            for col, val in enumerate(vals, 1):
                c           = _cell(ws, row_i, col)
                c.value     = _safe_text(val)
                c.font      = _FONT_SYS
                c.alignment = _ALIGN_L
                c.border    = _BORDER
                if row_i % 2 == 0:
                    c.fill = img_alt

        for col, w in {1: 14, 2: 60, 3: 20, 4: 20, 5: 22, 6: 18, 7: 18, 8: 14}.items():
            ws.column_dimensions[get_column_letter(col)].width = w
        ws.freeze_panes = "A20"

    # ── Reimport ──────────────────────────────────────────────────────────────

    def reimport_excel(self, excel_path: Path) -> dict:
        """
        Read an edited Excel and update the DB.
        Price recalculation uses apply_margin() with the persisted round_to / suffix
        from each product's DB row — single source of truth (C1 fix).
        """
        log.info(f"Re-importing Excel: {excel_path}")
        results = {
            "processed": 0, "updated": 0,
            "not_found": [], "duplicates": [], "errors": [], "warnings": [],
        }

        try:
            wb = load_workbook(str(excel_path), data_only=True)
        except Exception as e:
            results["errors"].append(f"Could not open Excel: {e}")
            return results

        ws = wb.worksheets[0]
        # Headers may live on row 1 (legacy) or row 2 (after staleness banner).
        # Detect by looking for "SKU" in the first two rows.
        header_row_idx = 1
        for candidate in (1, 2):
            row_vals = [c.value for c in ws[candidate]]
            if "SKU" in row_vals:
                header_row_idx = candidate
                break

        header_row = [c.value for c in ws[header_row_idx]]
        by_header  = {str(h).strip(): i + 1 for i, h in enumerate(header_row) if h}
        # Canonical key → column index, whichever alias the file uses
        col_map = {
            key: next((by_header[a] for a in aliases if a in by_header), None)
            for key, aliases in HEADER_ALIASES.items()
        }

        missing = [HEADER_ALIASES[k][0] for k in ("sku", "active") if col_map[k] is None]
        if missing:
            results["errors"].append(f"Missing required columns: {missing}")
            return results

        all_rows: list[dict] = []
        for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
            def _get(col_name, default=None):
                idx = col_map.get(col_name)
                return row[idx - 1] if idx else default

            sku = str(_get("sku") or "").strip().upper()
            if not sku:
                continue
            all_rows.append({
                "sku":         sku,
                "active":      str(_get("active", "YES") or "YES"),
                "brand":       _get("brand"),
                "description": _get("description"),
                "notes":       _get("notes"),
                "cost":        _get("cost"),
                "margin_pct":  _get("margin"),
                "iva_pct":     _get("tax"),
                "final_price": _get("final"),
                "category":    _get("category"),
            })

        # Dedup
        seen: dict[str, int] = {}
        for d in all_rows:
            seen[d["sku"]] = seen.get(d["sku"], 0) + 1
        dupes = [s for s, c in seen.items() if c > 1]
        if dupes:
            results["duplicates"] = dupes
            results["warnings"].append(
                f"Duplicate SKUs (first occurrence used): {dupes}"
            )
            unique: dict[str, dict] = {}
            for d in all_rows:
                if d["sku"] not in unique:
                    unique[d["sku"]] = d
            all_rows = list(unique.values())

        for d in all_rows:
            results["processed"] += 1
            sku       = d["sku"]
            is_active = d["active"].strip().lower() in TRUTHY

            existing = self.db.get_product(sku)
            if not existing:
                results["not_found"].append(sku)
                continue

            update_row: dict = {"sku": sku, "active": is_active}

            # Text fields
            for key, db_col in [
                ("brand",       "brand"),
                ("description", "description"),
                ("notes",       "notes"),
                ("category",    "category"),
            ]:
                val = d.get(key)
                if val is not None and str(val).strip():
                    update_row[db_col] = str(val).strip()

            def _to_float(v) -> Optional[float]:
                if v is None:
                    return None
                if isinstance(v, (int, float)):
                    return float(v)
                try:
                    return parse_number(v)
                except (ValueError, TypeError):
                    return None

            orig   = _to_float(d["cost"])
            margin = _to_float(d["margin_pct"])
            iva    = _to_float(d["iva_pct"])

            if orig   is not None: update_row["price_original"] = orig
            if margin is not None: update_row["margin_pct"]     = margin
            if iva    is not None: update_row["iva_pct"]        = iva

            # Recalculate price_final with apply_margin() using persisted round_to/suffix
            # This is the C1 fix: one code path for price calculation, always.
            from app.models.product import Product
            _orig   = orig   if orig   is not None else (existing.get("price_original") or 0.0)
            _margin = margin if margin is not None else (existing.get("margin_pct")     or 0.0)
            _iva    = iva    if iva    is not None else (existing.get("iva_pct")        or 0.0)
            _rnd    = existing.get("round_to_nearest") or 0
            _sfx    = existing.get("price_suffix") or ""

            p = Product(
                sku=sku, brand="", description="", price_original=_orig, category=""
            )
            p.apply_margin(_margin, _iva, _rnd, _sfx)
            update_row["price_final"]      = p.price_final
            update_row["round_to_nearest"] = p.round_to_nearest
            update_row["price_suffix"]     = p.price_suffix

            try:
                self.db.upsert_product(update_row)
                self.db.set_product_active(sku, is_active)
                results["updated"] += 1
            except Exception as e:
                results["errors"].append(f"Error updating {sku}: {e}")

        log.info(
            f"Reimport: processed={results['processed']} "
            f"updated={results['updated']} "
            f"not_found={len(results['not_found'])}"
        )
        return results
